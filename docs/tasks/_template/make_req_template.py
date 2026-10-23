# -*- coding: utf-8 -*-
"""生成《需求登记表》xlsx 模板（表单型）。改字段改本文件后重跑即可。

用法：
    python docs/tasks/_template/make_req_template.py [输出路径]
默认输出：docs/tasks/_template/req-template.xlsx
"""
import sys

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else r"D:/RuoyiProject/comic/docs/tasks/_template/req-template.xlsx"

BLUE = "FF2F5597"
WHITE = "FFFFFFFF"
ROWHEAD = "FFD9E2F3"
INPUT = "FFFFF9E6"
BORDER = "FFBFBFBF"
TIP = "FFF2F2F2"
RED = "FFC00000"
GRAY = "FF808080"
TEXT = "FF1A1A1A"

thin = Side(style="thin", color=BORDER)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FONT = "微软雅黑"

# (字段, 必填, 说明, 行高)
FIELDS = [
    ("任务标题", "是", "动词开头，≤20 字。示例：统一 check 脚本的双端覆盖", 26),
    ("需求类型", "是", "下拉选择：新功能 / 缺陷修复 / 重构 / 文档 / 运维 / 性能优化", 26),
    ("优先级", "是", "下拉选择：P0 紧急 / P1 高 / P2 中 / P3 低", 26),
    ("背景与问题", "是", "为什么做？现在有什么问题？可写日志或截图的文件名", 72),
    ("目标·完成态", "是", "做完之后能怎样？必须可验证，不要写「优化一下」", 72),
    ("涉及模块", "是", "从 9 个模块里选：comic-core / crawler-service / api-service / comic-scheduler / comic-web / comic-front / deploy / scripts / docs。不用写文件名", 44),
    ("改动范围（白名单）", "否", "选填。记得具体文件就写路径；不记得就留空 —— 设计阶段会按「涉及模块」细化到文件级后回填", 72),
    ("明确不做", "否", "防止范围膨胀。示例：不修 check.bat 缺 PYTHONPATH 的老问题", 72),
    ("是否涉及数据库变更", "是", "下拉：无 / 表结构变更 / 数据变更 / 两者都有。有变更时运维需迁移与备份", 36),
    ("验收标准", "是", "1~5 条，每条都要能跑命令或看到结果。一行一条", 72),
    ("参考材料", "否", "文件路径 / 截图文件名 / 链接", 40),
    ("期望完成时间", "否", "YYYY-MM-DD", 26),
    ("允许自动提交", "是", "下拉：否 / 是（默认否）。否 = 改完等你确认再提交", 26),
    ("补充说明", "否", "其他约束、偏好、注意事项", 72),
]

EXAMPLE = [
    "消息中心收藏漫画提示增加链接到详情页",
    "新功能",
    "P1 高",
    "现在消息中心当有收藏漫画更新时会提示出来，但需要自己手动找详情页，增加链接方便跳转。",
    "点击对应的收藏更新提示能跳转到对应详情页。",
    "comic-web",
    "（留空，由设计细化）",
    "不改后端；不改数据库；不处理返回逻辑",
    "无",
    "1) 点 update 条目跳到 /#/comic/<id>\n2) 无 comicId 的条目不跳、只标已读\n3) 两端 lint/typecheck/vitest 全绿",
    "scripts/NOTES.md",
    "2026-10-23",
    "否",
    "双端都要改",
]


def head_row(ws, row, cols, names):
    for col, name in zip(cols, names):
        c = ws[f"{col}{row}"]
        c.value = name
        c.font = Font(name=FONT, size=11, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = 26


wb = Workbook()

# ---------- Sheet 1: 需求单 ----------
ws = wb.active
ws.title = "需求单"

ws.merge_cells("A1:D1")
c = ws["A1"]
c.value = "comic 项目 · 需求登记表"
c.font = Font(name=FONT, size=14, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=BLUE)
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 36

ws.merge_cells("A2:D2")
c = ws["A2"]
c.value = "填写说明：浅黄格是填写区；标「必填」的请务必填；填完保存，把 xlsx 发回即可"
c.font = Font(name=FONT, size=11, color=GRAY)
c.fill = PatternFill("solid", fgColor=TIP)
c.alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[2].height = 24

head_row(ws, 3, "ABCD", ["字段", "填写内容", "必填", "说明与示例"])

for i, (name, req, desc, h) in enumerate(FIELDS):
    r = 4 + i
    a = ws[f"A{r}"]
    a.value = name
    a.font = Font(name=FONT, size=11, bold=True, color=TEXT)
    a.fill = PatternFill("solid", fgColor=ROWHEAD)
    a.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    a.border = BOX

    b = ws[f"B{r}"]
    b.fill = PatternFill("solid", fgColor=INPUT)
    b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    b.border = BOX
    b.font = Font(name=FONT, size=11, color=TEXT)

    cc = ws[f"C{r}"]
    cc.value = "必填" if req == "是" else "选填"
    cc.font = Font(name=FONT, size=10, bold=(req == "是"), color=RED if req == "是" else GRAY)
    cc.alignment = Alignment(horizontal="center", vertical="center")
    cc.border = BOX

    d = ws[f"D{r}"]
    d.value = desc
    d.font = Font(name=FONT, size=10, color=GRAY)
    d.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    d.border = BOX

    ws.row_dimensions[r].height = h

for col, w in zip("ABCD", [18, 54, 8, 50]):
    ws.column_dimensions[col].width = w

for cell, formula in [
    ("B5", '"新功能,缺陷修复,重构,文档,运维,性能优化"'),
    ("B6", '"P0 紧急,P1 高,P2 中,P3 低"'),
    ("B12", '"无,表结构变更,数据变更,两者都有"'),
    ("B16", '"否,是"'),
]:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(cell)

ws.freeze_panes = "A4"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = "A1:D17"

# ---------- Sheet 2: 填写说明 ----------
w2 = wb.create_sheet("填写说明")
w2.merge_cells("A1:D1")
c = w2["A1"]
c.value = "填写说明"
c.font = Font(name=FONT, size=13, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=BLUE)
c.alignment = Alignment(horizontal="center", vertical="center")
w2.row_dimensions[1].height = 32

RULES = [
    "1. 一行一个字段，不要把多条需求塞进一张表；多个需求请复制多份文件。",
    "2. 验收标准必须能跑命令或看到结果 —— 「能正常使用」这种不算。",
    "3. 「涉及模块」只选模块即可，不用记文件名；「改动范围」选填，记得就写，不记得留空。",
    "4. 不要写「优化一下 / 改改看 / 顺手修一下」，要写清楚改成什么。",
    "5. 有截图就另存一份，在「参考材料」里写文件名；不要把图片粘进单元格。",
    "6. 填完直接发 xlsx 本身（不要转 pdf、不要截图回传）。",
    "7. 涉及数据库变更的，在「补充说明」写清变更类型与执行顺序；运维会据此决定是否先备份。",
]

w2.merge_cells("A3:D3")
w2["A3"] = "怎么填"
w2["A3"].font = Font(name=FONT, size=11, bold=True, color=BLUE)
w2["A3"].fill = PatternFill("solid", fgColor=ROWHEAD)

for i, t in enumerate(RULES):
    r = 4 + i
    w2.merge_cells(f"A{r}:D{r}")
    c = w2[f"A{r}"]
    c.value = t
    c.font = Font(name=FONT, size=10, color=TEXT)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c.border = BOX
    w2.row_dimensions[r].height = 30

w2.merge_cells("A12:D12")
w2["A12"] = "常见错误"
w2["A12"].font = Font(name=FONT, size=11, bold=True, color=BLUE)
w2["A12"].fill = PatternFill("solid", fgColor=ROWHEAD)

head_row(w2, 13, "ABCD", ["错误写法", "正确写法", "", "为什么"])
w2.merge_cells("B13:C13")

MISTAKES = [
    ("优化一下首页加载", "首页首屏 LCP 从 3.2s 降到 1.5s 以内", "不可验证"),
    ("改改后端", "api-service/routers/public.py", "范围不清"),
    ("所有地方都查一下", "涉及模块选 comic-web（设计再定位到文件）", "范围膨胀"),
    ("验收：能正常使用", "验收：访问 /#/ 返回 200，列表渲染 ≥ 10 条", "无法判定"),
    ("（不写不做的事）", "不做：不改数据库结构、不动 deploy", "容易越界"),
]
for i, (bad, good, why) in enumerate(MISTAKES):
    r = 14 + i
    w2[f"A{r}"] = bad
    w2[f"A{r}"].font = Font(name=FONT, size=10, color=RED)
    w2[f"A{r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    w2[f"A{r}"].border = BOX
    w2.merge_cells(f"B{r}:C{r}")
    w2[f"B{r}"] = good
    w2[f"B{r}"].font = Font(name=FONT, size=10, color=TEXT)
    w2[f"B{r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    w2[f"B{r}"].border = BOX
    w2[f"C{r}"].border = BOX
    w2[f"D{r}"] = why
    w2[f"D{r}"].font = Font(name=FONT, size=10, color=GRAY)
    w2[f"D{r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    w2[f"D{r}"].border = BOX
    w2.row_dimensions[r].height = 30

w2.merge_cells("A20:D20")
w2["A20"] = "回传后会发生什么"
w2["A20"].font = Font(name=FONT, size=11, bold=True, color=BLUE)
w2["A20"].fill = PatternFill("solid", fgColor=ROWHEAD)

STEPS = [
    "第 1 步：解析这张表 → 生成 01-brief.md（任务卡：范围白名单 + 验收标准）",
    "第 2 步：你在对话里确认一次（也可以直接说「开」，跳过等待）",
    "第 3 步：全自动跑完 设计 → 开发 → 测试 → 收尾；门禁不过自动打回上一阶段",
]
for i, t in enumerate(STEPS):
    r = 21 + i
    w2.merge_cells(f"A{r}:D{r}")
    c = w2[f"A{r}"]
    c.value = t
    c.font = Font(name=FONT, size=10, color=TEXT)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c.border = BOX
    w2.row_dimensions[r].height = 26

for col, w in zip("ABCD", [30, 40, 22, 26]):
    w2.column_dimensions[col].width = w

# ---------- Sheet 3: 示例 ----------
w3 = wb.create_sheet("示例")
w3.merge_cells("A1:D1")
c = w3["A1"]
c.value = "示例：一张填好的需求单（照这个粒度填）"
c.font = Font(name=FONT, size=13, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=BLUE)
c.alignment = Alignment(horizontal="center", vertical="center")
w3.row_dimensions[1].height = 34

w3.merge_cells("A2:D2")
c = w3["A2"]
c.value = "行位与「需求单」完全一致（表头第 3 行、数据第 4~17 行），可直接对照填写"
c.font = Font(name=FONT, size=10, color=GRAY)
c.fill = PatternFill("solid", fgColor=TIP)
c.alignment = Alignment(horizontal="left", vertical="center")
w3.row_dimensions[2].height = 22

head_row(w3, 3, "ABCD", ["字段", "填写内容", "必填", "说明"])

for i, (name, req, desc, h) in enumerate(FIELDS):
    r = 4 + i
    a = w3[f"A{r}"]
    a.value = name
    a.font = Font(name=FONT, size=11, bold=True, color=TEXT)
    a.fill = PatternFill("solid", fgColor=ROWHEAD)
    a.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    a.border = BOX

    b = w3[f"B{r}"]
    b.value = EXAMPLE[i]
    b.font = Font(name=FONT, size=10, color=TEXT)
    b.fill = PatternFill("solid", fgColor=INPUT)
    b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    b.border = BOX

    cc = w3[f"C{r}"]
    cc.value = req
    cc.font = Font(name=FONT, size=10, color=GRAY)
    cc.alignment = Alignment(horizontal="center", vertical="center")
    cc.border = BOX

    d = w3[f"D{r}"]
    d.value = desc
    d.font = Font(name=FONT, size=10, color=GRAY)
    d.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    d.border = BOX

    w3.row_dimensions[r].height = max(h, 36)

for col, w in zip("ABCD", [18, 54, 8, 50]):
    w3.column_dimensions[col].width = w
w3.freeze_panes = "A3"

wb.properties.title = "需求登记表"
wb.save(OUT)
print(f"OK -> {OUT}")
print(f"sheets: {wb.sheetnames}")
